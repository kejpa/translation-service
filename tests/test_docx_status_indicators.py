from docx import Document

from translation_service.docx_status_indicators import (
    apply_left_border,
    apply_status_indicator,
    get_border_color,
)
from translation_service.translation_status import TranslationStatus


def test_translated_has_no_indicator():
    assert (
        get_border_color(
            TranslationStatus.TRANSLATED,
        )
        is None
    )


def test_empty_has_no_indicator():
    assert (
        get_border_color(
            TranslationStatus.EMPTY,
        )
        is None
    )


def test_fuzzy_high_uses_green_indicator():
    assert (
        get_border_color(
            TranslationStatus.FUZZY_HIGH,
        )
        == "00FF00"
    )


def test_fuzzy_low_uses_yellow_indicator():
    assert (
        get_border_color(
            TranslationStatus.FUZZY_LOW,
        )
        == "FFFF00"
    )


def test_llm_uses_red_indicator():
    assert (
        get_border_color(
            TranslationStatus.LLM,
        )
        == "FF0000"
    )


def test_missing_uses_red_indicator():
    assert (
        get_border_color(
            TranslationStatus.MISSING,
        )
        == "FF0000"
    )


def test_apply_left_border_adds_border_color_to_xml():
    document = Document()

    paragraph = document.add_paragraph(
        "Test",
    )

    apply_left_border(
        paragraph,
        "00FF00",
    )

    assert "00FF00" in paragraph._p.xml


def test_fuzzy_high_applies_green_border():
    document = Document()

    paragraph = document.add_paragraph(
        "Test",
    )

    apply_status_indicator(
        paragraph,
        TranslationStatus.FUZZY_HIGH,
    )

    assert "00FF00" in paragraph._p.xml


def test_fuzzy_low_applies_yellow_border():
    document = Document()

    paragraph = document.add_paragraph(
        "Test",
    )

    apply_status_indicator(
        paragraph,
        TranslationStatus.FUZZY_LOW,
    )

    assert "FFFF00" in paragraph._p.xml


def test_llm_applies_red_border():
    document = Document()

    paragraph = document.add_paragraph(
        "Test",
    )

    apply_status_indicator(
        paragraph,
        TranslationStatus.LLM,
    )

    assert "FF0000" in paragraph._p.xml
