from starlette.testclient import TestClient

from tests.helpers import add_translation, create_docx
from translation_service.main import app
from io import BytesIO

from docx import Document

client = TestClient(app)


def test_download_returns_docx(
    db,
):
    add_translation(
        db,
        "Hei maailma",
        "Hej världen",
    )

    translate_response = client.post(
        "/docx/translate",
        files={
            "file": (
                "source.docx",
                create_docx("Hei maailma"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        },
    )

    assert translate_response.status_code == 200

    download_url = translate_response.json()["download_url"]

    response = client.get(
        download_url,
    )

    assert response.status_code == 200

    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    assert "attachment" in response.headers["content-disposition"]


def test_download_returns_translated_content(
    db,
):
    add_translation(
        db,
        "Hei maailma",
        "Hej världen",
    )

    translate_response = client.post(
        "/docx/translate",
        files={
            "file": (
                "source.docx",
                create_docx(
                    "Hei maailma",
                    "Tuntematon teksti",
                ),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        },
    )

    download_url = translate_response.json()["download_url"]

    response = client.get(
        download_url,
    )

    document = Document(BytesIO(response.content))

    paragraphs = [paragraph.text for paragraph in document.paragraphs]

    assert paragraphs == [
        "Hej världen",
        "Tuntematon teksti",
    ]


def test_download_preserves_empty_paragraphs(
    db,
):
    add_translation(
        db,
        "Hei maailma",
        "Hej världen",
    )

    add_translation(
        db,
        "Miten voit?",
        "Hur mår du?",
    )

    translate_response = client.post(
        "/docx/translate",
        files={
            "file": (
                "source.docx",
                create_docx(
                    "Hei maailma",
                    "",
                    "Miten voit?",
                ),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        },
    )

    download_url = translate_response.json()["download_url"]

    response = client.get(
        download_url,
    )

    document = Document(BytesIO(response.content))

    paragraphs = [paragraph.text for paragraph in document.paragraphs]

    assert paragraphs == [
        "Hej världen",
        "",
        "Hur mår du?",
    ]


def test_translate_docx_uses_requested_filename(
    db,
):
    add_translation(
        db,
        "Hei maailma",
        "Hej världen",
    )

    response = client.post(
        "/docx/translate",
        data={
            "output_filename": "my-translation.docx",
        },
        files={
            "file": (
                "source.docx",
                create_docx("Hei maailma"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        },
    )

    assert response.status_code == 200

    assert response.json()["download_url"] == ("/downloads/my-translation.docx")


def test_translate_docx_uses_default_filename(
    db,
):
    add_translation(
        db,
        "Hei maailma",
        "Hej världen",
    )

    response = client.post(
        "/docx/translate",
        files={
            "file": (
                "source.docx",
                create_docx("Hei maailma"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        },
    )

    assert response.status_code == 200

    assert response.json()["download_url"] == ("/downloads/translated.docx")
